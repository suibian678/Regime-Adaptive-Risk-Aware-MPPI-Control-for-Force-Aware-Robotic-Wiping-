#!/usr/bin/env python3
"""Build the final Chinese adviser-response Word report from frozen artefacts."""

from __future__ import annotations

from copy import deepcopy
from datetime import date
import json
from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT.parent / "ForceWipe_论文导师意见修改说明_2026-09-14.docx"
ANALYSIS = ROOT / "results/analysis/deployment_gap_stress_20260920/RESULT.json"
AUDIT = ROOT / "results/analysis/deployment_gap_stress_20260920/INDEPENDENT_AUDIT.json"
NUMBER_AUDIT = (
    ROOT / "results/analysis/manuscript_experiment_design_audit/"
    "MANUSCRIPT_EXPERIMENT_NUMBER_AUDIT.json"
)
CROSSSIM_RESULT = (
    ROOT / "results/crosssim/v19_mujoco_zero_shot_transfer_20260921_r1/RESULT.json"
)
CROSSSIM_AUDIT = (
    ROOT / "results/analysis/v19_mujoco_zero_shot_transfer_independent_audit.json"
)
OUTPUT = ROOT / "ForceWipe_导师意见修改说明_方法与扩展实验_2026-09-21.docx"

BLUE = "1F4E78"
PALE_BLUE = "D9EAF7"
GRID = "C8D0D9"
GREEN = RGBColor(46, 125, 50)
AMBER = RGBColor(183, 121, 31)
GREY = RGBColor(90, 90, 90)


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    element = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        child = OxmlElement(f"w:{edge}")
        child.set(qn("w:val"), "single")
        child.set(qn("w:sz"), "4")
        child.set(qn("w:space"), "0")
        child.set(qn("w:color"), GRID)
        element.append(child)
    tbl_pr.append(element)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_margin(cell, top=70, start=85, bottom=70, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_run(run, size=9.0, bold=False, colour=None) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(size)
    run.bold = bold
    if colour is not None:
        run.font.color.rgb = colour


def set_cell_text(cell, text: str, *, header=False, centre=False, status=False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centre else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(str(text))
    colour = RGBColor(255, 255, 255) if header else None
    if status and str(text) in {"已完成", "已完成至仿真边界", "通过"}:
        colour = GREEN
    elif status:
        colour = AMBER
    format_run(run, size=9.0, bold=header or status, colour=colour)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margin(cell)


def add_table(doc, headers, rows, widths, centre_columns=(), status_column=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = Inches(width)
        shade(cell, BLUE)
        set_cell_text(cell, header, header=True, centre=True)
    repeat_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        keep_row(table.rows[-1])
        if row_index % 2 == 0:
            for cell in cells:
                shade(cell, PALE_BLUE)
        for column, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Inches(width)
            set_cell_text(
                cell, value,
                centre=column in centre_columns,
                status=(status_column is not None and column == status_column),
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    format_run(run, size=14, bold=True)


def add_body(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.18
    run = paragraph.add_run(text)
    format_run(run, size=10)


def deployment_summary(analysis: dict) -> tuple[str, str]:
    rows = analysis["condition_summaries"]
    endpoint_ids = {
        "force_noise_rms_0p035n",
        "force_bias_m0p8n",
        "force_bias_p0p8n",
        "delay_5samples",
        "tcp_pose_error_0p1mm",
        "normal_error_5deg",
    }
    endpoints = [row for row in rows if row["condition_id"] in endpoint_ids]
    nominal = [row for row in rows if row["condition_id"] == "nominal"]
    endpoint_compound = sum(int(row["compound_passes"]) for row in endpoints)
    endpoint_evaluations = sum(int(row["evaluations"]) for row in endpoints)
    nominal_compound = sum(int(row["compound_passes"]) for row in nominal)
    nominal_evaluations = sum(int(row["evaluations"]) for row in nominal)
    worst = min(endpoints, key=lambda row: (row["compound_passes"], -row["mean_mre"]))
    overview = (
        f"完成 375 次新增压力评估并复用 45 次 nominal，共报告 420 次；"
        f"六个最大端点合计 {endpoint_compound}/{endpoint_evaluations} compound pass，"
        f"对应 nominal 为 {nominal_compound}/{nominal_evaluations}。"
    )
    detail = (
        f"最弱端点为 {worst['block_id']}/{worst['condition_id']}："
        f"{worst['compound_passes']}/{worst['evaluations']} compound pass，"
        f"mean MRE {100 * worst['mean_mre']:.2f}%，"
        f"最大峰值 {worst['maximum_peak_force_n']:.3f} N。"
    )
    return overview, detail


def main() -> None:
    if not REFERENCE.is_file():
        raise SystemExit(f"reference missing: {REFERENCE}")
    if (not ANALYSIS.is_file() or not AUDIT.is_file() or not NUMBER_AUDIT.is_file()
            or not CROSSSIM_RESULT.is_file() or not CROSSSIM_AUDIT.is_file()):
        raise SystemExit("completed analysis and audits are required before Word generation")
    analysis = load(ANALYSIS)
    audit = load(AUDIT)
    number_audit = load(NUMBER_AUDIT)
    crosssim_result = load(CROSSSIM_RESULT)
    crosssim_audit = load(CROSSSIM_AUDIT)
    if analysis.get("status") != "completed" or audit.get("status") != "PASS":
        raise SystemExit("deployment-gap artefacts are not complete")
    if crosssim_result.get("status") != "completed" or crosssim_audit.get("status") != "PASS":
        raise SystemExit("cross-engine artefacts are not complete")
    gap_overview, gap_detail = deployment_summary(analysis)

    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.left_margin = Inches(0.50)
    section.right_margin = Inches(0.50)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(10)

    title_style_pr = doc.styles["Title"]._element.get_or_add_pPr()
    title_style_border = title_style_pr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_pr.remove(title_style_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    title_pr = title._p.get_or_add_pPr()
    title_border = title_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_pr.remove(title_border)
    format_run(title.add_run("ForceWipe 方法论文导师意见修改说明"), size=26, bold=True)
    date_line = doc.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_line.paragraph_format.space_after = Pt(8)
    format_run(date_line.add_run("2026 年 9 月 21 日"), size=10, colour=GREY)

    add_body(
        doc,
        "本轮按照导师关于“将评测型论文升级为方法型论文”的意见，完成了方法重构、配套消融、"
        "因子分离评估、公平性补充实验及外部有效性检查。论文的核心主张调整为：直接 TD-MPC2 "
        "能够完成多目标力擦拭，提出的风险感知、力条件化规划机制进一步改善力跟踪与计算开销之间的权衡。"
        "本文仍为仿真研究，不将仿真压力测试或跨仿真器校准写成真机验证。"
    )

    add_heading(doc, "一  完成情况概览")
    overview_rows = [
        ("方法升级", "已完成", "力条件化潜在动力学、下一步力预测、连续状态权重及自适应 MPPI 计算", "形成可独立描述和消融的方法贡献"),
        ("方法消融", "已完成", "目标自适应、计算自适应、风险表征、BC 强度和历史路由阈值", "识别有效组成并删除无证据支持的复杂机制"),
        ("因子分离评估", "已完成", "表面、路径、摩擦和支撑刚度分别变化；270 次配对评估", "避免多因素同时变化造成归因混淆"),
        ("PPO 公平性补充", "已完成", "3 个训练预算、6 个局部超参数变体、5 个种子及 225 次冻结确认", "回应同预算不等于同等优化质量"),
        ("跨仿真器验证", "已完成", "完成 policy-independent 校准，并将 5 个冻结 checkpoint 零样本部署到 MuJoCo 的平面、倾斜面和圆柱面", "平面与倾斜面 30/30 完整通过；圆柱面 0/15，明确暴露曲面迁移边界"),
        ("部署差距压力测试", "已完成", gap_overview, "补充真机缺失条件下的可部署性边界证据"),
        ("真机实验", "未完成", "当前无可用硬件平台", "作为明确限制保留"),
    ]
    add_table(
        doc, ("模块", "状态", "完成内容", "论文作用"), overview_rows,
        (1.35, 0.85, 5.35, 2.45), centre_columns=(0, 1), status_column=1,
    )

    add_heading(doc, "二  主要意见逐项修改")
    response_rows = [
        ("1", "从 TD-MPC2+BC+MPPI 路由升级为风险感知、力条件化控制", "目标力在潜在转移中持续保留；学习的下一步力预测形成因果风险量；代价与预算随接触状态连续变化", "方法节统一公式与算法流程", "已完成"),
        ("2", "0.85 人工路由阈值创新性弱且缺乏敏感性", "新方法删除硬切换；旧方法在 0.70–0.95 范围进行冻结敏感性复算", "阈值改变 MPPI 使用比例，但未形成自然最优点", "已完成"),
        ("3", "证明 MPPI 的作用并报告计算代价", "同一 checkpoint 比较固定目标/计算、自适应目标、自适应计算和完整方法", "compound pass 85/135→100/135；平均少 7.97 个样本，规划中位耗时少 25.65 ms", "已完成"),
        ("4", "补充 BC 消融", "在相同数据、更新次数、种子和规划器下比较 BC=0、0.5、2", "无 BC 时 15/15 无有效接触；强 BC 保持全部种子可用", "已完成"),
        ("5", "OOD 多因素同时变化导致归因困难", "建立 reference、倾斜面、圆柱、两条路径、两档摩擦和两档刚度的单因素设计", "270 次配对评估；圆柱 12 N 仍是主要边界", "已完成"),
        ("6", "PPO 预算和超参数不足，比较可能不公平", "新增 1×/2×/4× 预算及学习率、clip ratio、entropy coefficient 局部扫描；DEV 选择后在冻结套件确认", "最佳配置为 clip 0.1、2×预算；确认集 6/225 compound pass，且集中于一个种子", "已完成"),
        ("7", "Adaptive admittance 重复次数不对称", "明确其为每单元一次的确定性结构控制参考，不进入跨种子配对统计", "正文不再将其表述为统计对称比较", "已完成"),
        ("8", "12 N 欠跟踪应发展为 safety–tracking trade-off", "显式分离 tracking 与 transient risk，并按状态连续调整权重和计算", "12 N tracking 5/45→17/45；任务与超限未同步改善，贡献限定为 tracking–computation 权衡", "已完成"),
        ("9", "direct 不能被误读为端到端感知或无监督训练", "将 direct 定义为部署动作无经典力控、shield 或补擦；同步披露 BC 与注册名义坐标系", "摘要、任务定义与局限性使用同一口径", "已完成"),
        ("10", "跨仿真器及更完整 Sim2Real 证据", "先记录 MuJoCo 端口校准误差，再零样本运行 45 个匹配 checkpoint–surface–force 单元；另完成 force/latency/TCP/frame 单因素压力测试", f"MuJoCo 平面与倾斜面 30/30 compound pass，圆柱面 0/15；保留的切向位置误差使该结果定位为跨引擎压力测试。{gap_detail}", "已完成至仿真边界"),
    ]
    add_table(
        doc, ("序号", "导师意见", "已完成修改", "结果或正文证据", "状态"), response_rows,
        (0.48, 2.15, 3.55, 3.00, 0.82), centre_columns=(0, 4), status_column=4,
    )

    add_heading(doc, "三  语言句法与术语修改")
    language_rows = [
        ("1", "方法主张过宽", "将主张限定为仿真中的直接学习擦拭与 tracking–computation 改善，不声称一般安全提升", "摘要、引言、讨论、结论", "已完成"),
        ("2", "缩写首次出现未展开", "首次出现时展开 TD-MPC2、MPPI、BC、MRE、NRMSE 和 OOD", "全文", "已完成"),
        ("3", "结果段混入过多范围辩护", "结果段优先报告问题、效应量和区间；证据边界集中于设计与讨论", "实验设计、结果、讨论", "已完成"),
        ("4", "100 Hz 容易被误写成实时规划", "明确 100 Hz 是控制/仿真离散率，不代表 MPPI 达到 100 Hz 墙钟实时", "讨论", "已完成"),
        ("5", "cross-simulator 容易被写成 Sim2Real", "将 45 次迁移定义为 zero-shot cross-engine stress test，并同步报告校准残差", "实验设计、结果、讨论", "已完成"),
        ("6", "全文语言一致性", "统一为自然英式学术英语；保留算法名、引文标题和标准术语", "全文", "已完成"),
    ]
    add_table(
        doc, ("序号", "问题", "修改", "位置", "状态"), language_rows,
        (0.48, 2.10, 4.50, 1.92, 1.00), centre_columns=(0, 3, 4), status_column=4,
    )

    add_heading(doc, "四  最终核验与剩余事项")
    verification_rows = [
        ("Deployment-gap 独立复算", f"{audit['checks_passed']}/{audit['checks_total']} PASS；375 条 trace 的哈希、主键、seed、指标与权限逐项复算", "通过"),
        ("论文数字审计", f"{number_audit['checks_passed']}/{number_audit['checks_total']} PASS；覆盖主结果、PPO、跨引擎和 deployment-gap", "通过"),
        ("跨仿真器独立复算", f"{crosssim_audit['checks_passed']}/{crosssim_audit['checks_total']} PASS；45 条轨迹的主键、哈希、指标、planner seed 与匹配关系闭合", "通过"),
        ("引用与缩写", "所有引用键均存在；正文首次出现展开主要缩写；官方硬件规格单独引用", "通过"),
        ("PDF 版面", "最终 PDF 逐页渲染检查图表、表格、边界、字号与交叉引用", "通过"),
        ("真机与真实传感链", "尚未完成真实视觉、力传感、通信和硬件实时性验证；论文明确列为后续工作", "未完成"),
        ("标题与匿名归档", "标题继续留空等待导师确认；匿名代码与数据链接在投稿前替换占位符", "投稿前处理"),
    ]
    add_table(
        doc, ("核验项", "结论", "状态"), verification_rows,
        (2.05, 6.90, 1.05), centre_columns=(0, 2), status_column=2,
    )

    doc.core_properties.title = "ForceWipe 方法论文导师意见修改说明"
    doc.core_properties.subject = "方法升级、扩展实验与终态核验"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
